# -*- coding: utf-8 -*-
"""
日切 PKL（mongo_daily_deals_dump 格式）按合约统计「剔除做市/模拟金后」的每日日末净持仓 USDT，
并导出宽表与每 symbol 折线图。

口径：
- 每行一笔用户腿；ensure_user_id 后若 user_id 属于 exclude_user_ids，整行丢弃。
- 多头 long_coin >= 0；空头 short_coin <= 0（越负空头越大）；净币量 net_coin = long_coin + short_coin。
- 日末净 USDT = 日末 net_coin × 该 symbol 最近一笔成交价（当日无成交则沿用上一有效价）。

依赖：pandas、numpy、xlwt（.xls）或 openpyxl（.xlsx）、matplotlib；Python 3.9+。

PKL 目录示例（Ubuntu）：
  /home/ubuntu/strategy_base/strategy/openclaw_analysis_user/data/mongo_daily_deals_pkl

示例：
  python pkl_symbol_net_position_daily.py --data-dir ./data/mongo_daily_deals_pkl

默认写出 net_position.xls，图表目录 ./chart；仅写 xlsx 时可传 --output-xls "" --output-xlsx out.xlsx。
xlsx 在主表下方另附：max(正持仓均值,|负持仓均值|) 超过 10000 USDT 的交易对（一行一个）。
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import DefaultDict, Dict, Iterable, List, Optional, Set, Tuple

import numpy as np
import pandas as pd

_DIR = os.path.dirname(os.path.abspath(__file__))
if _DIR not in sys.path:
    sys.path.insert(0, _DIR)

from exclude_user_ids import get_excluded_user_ids_sync  # noqa: E402

DEFAULT_DATA_DIR = os.path.join(_DIR, "data", "mongo_daily_deals_pkl")
UBUNTU_PKL_HINT = (
    "/home/ubuntu/strategy_base/strategy/openclaw_analysis_user/data/mongo_daily_deals_pkl"
)

XLS_MAX_COLS = 256
INVALID_FN_CHARS = '\\/:*?"<>|'

POSITION_STAT_COLS = (
    "最小持仓",
    "最大持仓",
    "正持仓均值",
    "正持仓中位数",
    "负持仓均值",
    "负持仓中位数",
)

# xlsx 底部附录：max(正持仓均值, |负持仓均值|) 超过该阈值的交易对（与正/负均值的 nan 口径一致：缺一侧按 0 参与 max）
EXTREME_MEAN_THRESHOLD_U = 10000.0
XLSX_APPEND_HEADER_SYMBOL = "交易对"
XLSX_APPEND_HEADER_VALUE = "max(正持仓均值,|负持仓均值|)"

BUY_SELL_CN = {"开多", "开空", "平多", "平空"}
NUM_SIDE = {"1": "开多", "2": "开空", "3": "平多", "4": "平空"}


def ensure_user_id(df: pd.DataFrame) -> pd.DataFrame:
    """与 pkl_target_users_trade_stats.ensure_user_id 一致。"""
    d = df.copy()
    if "user_id" in d.columns:
        d["user_id"] = d["user_id"].astype(str)
        return d
    if "taker_user" not in d.columns or "maker_user" not in d.columns:
        raise ValueError("pkl 缺少 user_id / taker_user / maker_user，无法恢复 user_id")

    d = d.reset_index(drop=True)
    grp_cols = ["ts_text", "symbol", "price", "amount", "taker_user", "maker_user"]
    for c in grp_cols:
        if c not in d.columns:
            raise ValueError("pkl 缺少字段 {}，无法恢复 user_id".format(c))
    d["_idx_in_pair"] = d.groupby(grp_cols).cumcount()
    d["user_id"] = d["taker_user"].astype(str)
    d.loc[(d["_idx_in_pair"] % 2) == 1, "user_id"] = d.loc[
        (d["_idx_in_pair"] % 2) == 1, "maker_user"
    ].astype(str)
    d = d.drop(columns=["_idx_in_pair"])
    return d


def list_pkl_paths(data_dir: str) -> List[str]:
    if not os.path.isdir(data_dir):
        raise FileNotFoundError("PKL 目录不存在: {}".format(data_dir))
    out: List[str] = []
    for fn in sorted(os.listdir(data_dir)):
        if not fn.endswith(".pkl"):
            continue
        stem = fn[:-4]
        try:
            datetime.strptime(stem, "%Y-%m-%d")
        except ValueError:
            continue
        out.append(os.path.join(data_dir, fn))
    return out


def load_concat_pkls(paths: Iterable[str]) -> pd.DataFrame:
    ps = list(paths)
    if not ps:
        return pd.DataFrame()
    dfs = [pd.read_pickle(p) for p in ps]
    return pd.concat(dfs, ignore_index=True)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="PKL 剔除做市/模拟金后，按 symbol 日末净持仓 USDT 宽表 + 折线图"
    )
    p.add_argument(
        "--data-dir",
        type=str,
        default=os.environ.get("UPM_PKL_DEALS_DIR", DEFAULT_DATA_DIR),
        help="日切 PKL 目录（YYYY-MM-DD.pkl）；Ubuntu 可参考: {}".format(UBUNTU_PKL_HINT),
    )
    p.add_argument(
        "--output-xls",
        type=str,
        default="net_position.xls",
        help="输出 .xls 宽表路径（默认 net_position.xls；不需要时传空串；列数含 6 列统计列不得超过 256，否则请用 --output-xlsx）",
    )
    p.add_argument(
        "--output-xlsx",
        type=str,
        default="",
        help="输出 .xlsx 宽表（列数不受 256 限制）",
    )
    p.add_argument(
        "--charts-dir",
        type=str,
        default="./chart",
        help="每个 symbol 写一张 PNG 折线图（默认 ./chart；不需要图表时传空串）",
    )
    p.add_argument(
        "--start-col-date",
        type=str,
        default="2026-02-01",
        help="宽表起始列日期 YYYY-MM-DD（默认 2026-02-01）",
    )
    p.add_argument(
        "--no-exclude",
        action="store_true",
        help="调试：不做做市/模拟金剔除（勿用于生产口径）",
    )
    args, _ = p.parse_known_args()
    return args


def _normalize_buy_sell(raw: object) -> Optional[str]:
    if raw is None or (isinstance(raw, float) and np.isnan(raw)):
        return None
    s = str(raw).strip()
    if s in BUY_SELL_CN:
        return s
    if s in NUM_SIDE:
        return NUM_SIDE[s]
    return None


def _day_range(start_d: date, end_d: date) -> List[date]:
    out: List[date] = []
    cur = start_d
    while cur <= end_d:
        out.append(cur)
        cur += timedelta(days=1)
    return out


def safe_chart_basename(symbol: str, ext: str = ".png") -> str:
    s = str(symbol).strip() or "unknown"
    for ch in INVALID_FN_CHARS:
        s = s.replace(ch, "_")
    if len(s) > 180:
        s = s[:180]
    return s + ext


def simulate_eod_matrix(
    df: pd.DataFrame,
    start_col_d: date,
) -> Tuple[List[date], List[str], Dict[Tuple[str, date], float], int, date, date]:
    """
    返回 output_dates, symbols_sorted, eod[(symbol, d)], skipped, min_d, max_d。
    eod 仅对 [min_d, max_d] 内每个自然日、每个 symbol 有记录；更早的输出列由调用方填 0。
    """
    if df.shape[0] == 0:
        return [], [], {}, 0, start_col_d, start_col_d

    work = df.copy()
    work["_ts"] = pd.to_datetime(work["ts_text"], errors="coerce")
    work = work[work["_ts"].notna()].copy()
    if work.shape[0] == 0:
        return [], [], {}, 0, start_col_d, start_col_d

    work["_d"] = work["_ts"].dt.date
    work = work.sort_values(["_d", "ts_text"], kind="mergesort").reset_index(drop=True)

    all_symbols: Set[str] = set(work["symbol"].astype(str))
    min_d = work["_d"].min()
    max_d = work["_d"].max()

    state_long: DefaultDict[str, float] = defaultdict(float)
    state_short: DefaultDict[str, float] = defaultdict(float)
    last_px: Dict[str, float] = {}

    eod: Dict[Tuple[str, date], float] = {}
    skipped = 0

    by_day: Dict[date, pd.DataFrame] = {k: v for k, v in work.groupby("_d", sort=True)}

    for d in _day_range(min_d, max_d):
        day = by_day.get(d)
        if day is not None and day.shape[0] > 0:
            day = day.sort_values("ts_text", kind="mergesort")
            for i in range(day.shape[0]):
                row = day.iloc[i]
                sym = str(row["symbol"])
                side = _normalize_buy_sell(row.get("buy_sell"))
                pr = float(pd.to_numeric(row.get("price"), errors="coerce") or 0.0)
                amt = float(pd.to_numeric(row.get("amount"), errors="coerce") or 0.0)
                fv = float(pd.to_numeric(row.get("face_value"), errors="coerce") or 0.0)
                if fv <= 0:
                    fv = 1.0
                coin_qty = amt * fv

                if side is None:
                    skipped += 1
                    continue

                lc = state_long[sym]
                sc = state_short[sym]

                if side == "开多":
                    lc += coin_qty
                elif side == "平多":
                    lc -= min(coin_qty, lc)
                elif side == "开空":
                    sc -= coin_qty
                elif side == "平空":
                    sc += min(coin_qty, abs(sc))
                else:
                    skipped += 1
                    continue

                state_long[sym] = lc
                state_short[sym] = sc
                if pr > 0:
                    last_px[sym] = pr

        for sym in all_symbols:
            px = last_px.get(sym)
            net_coin = state_long[sym] + state_short[sym]
            if px is None or px <= 0:
                eod[(sym, d)] = 0.0
            else:
                eod[(sym, d)] = net_coin * px

    out_dates = _day_range(start_col_d, max_d) if max_d >= start_col_d else []
    syms_sorted = sorted(all_symbols)
    return out_dates, syms_sorted, eod, skipped, min_d, max_d


def _cell_net_usdt(
    sym: str,
    d: date,
    eod: Dict[Tuple[str, date], float],
    min_d: date,
    max_d: date,
) -> float:
    if d < min_d or d > max_d:
        return 0.0
    return float(eod.get((sym, d), 0.0))


def _xls_stat_cell(v: float) -> object:
    """xlwt 对 nan 不友好，无样本时写空单元格。"""
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return ""
    return float(v)


def _symbol_position_stats(
    sym: str,
    dates: List[date],
    eod: Dict[Tuple[str, date], float],
    min_d: date,
    max_d: date,
) -> Tuple[float, float, float, float, float, float]:
    """宽表日期列上日末净持仓：最小/最大（含 0）；正/负子集的均值与中位数（净持仓=0 的日期不参与）。"""
    vals = [_cell_net_usdt(sym, d, eod, min_d, max_d) for d in dates]
    arr = np.array(vals, dtype=float)
    mask = np.isfinite(arr)
    nan6 = (0.0, 0.0, float("nan"), float("nan"), float("nan"), float("nan"))
    if not mask.any():
        return nan6
    sub = arr[mask]
    pos = sub[sub > 0.0]
    neg = sub[sub < 0.0]
    m_pos = float(np.mean(pos)) if pos.size else float("nan")
    med_pos = float(np.median(pos)) if pos.size else float("nan")
    m_neg = float(np.mean(neg)) if neg.size else float("nan")
    med_neg = float(np.median(neg)) if neg.size else float("nan")
    return (
        float(np.min(sub)),
        float(np.max(sub)),
        m_pos,
        med_pos,
        m_neg,
        med_neg,
    )


def _max_pos_mean_or_abs_neg_mean(mpos: float, mneg: float) -> float:
    """正持仓均值为 nan 时按 0；负持仓均值为 nan 时按 0；负侧取绝对值后与正侧取 max。"""
    a = float(mpos) if np.isfinite(mpos) else 0.0
    b = abs(float(mneg)) if np.isfinite(mneg) else 0.0
    return float(max(a, b))


def _append_xlsx_extreme_mean_block(path: str, rows: List[Tuple[str, float]]) -> None:
    """在已存在的 xlsx 主表下方空一行后写入附录表头 + 一行一个交易对。"""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    r0 = ws.max_row + 2
    ws.cell(row=r0, column=1, value=XLSX_APPEND_HEADER_SYMBOL)
    ws.cell(row=r0, column=2, value=XLSX_APPEND_HEADER_VALUE)
    r = r0 + 1
    for sym, val in rows:
        ws.cell(row=r, column=1, value=sym)
        ws.cell(row=r, column=2, value=float(val))
        r += 1
    wb.save(path)


def write_matrix_xls(
    path: str,
    dates: List[date],
    symbols: List[str],
    eod: Dict[Tuple[str, date], float],
    min_d: date,
    max_d: date,
) -> None:
    try:
        import xlwt
    except ImportError as e:
        raise ImportError("导出 .xls 需要安装 xlwt：pip install xlwt") from e

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("net_usdt")

    ws.write(0, 0, "symbol")
    for j, d in enumerate(dates):
        ws.write(0, j + 1, d.isoformat())

    base_stat = 1 + len(dates)
    for k, name in enumerate(POSITION_STAT_COLS):
        ws.write(0, base_stat + k, name)

    for i, sym in enumerate(symbols):
        ws.write(i + 1, 0, sym)
        for j, d in enumerate(dates):
            v = _cell_net_usdt(sym, d, eod, min_d, max_d)
            ws.write(i + 1, j + 1, float(v))
        mn, mx, mpos, mdpos, mneg, mdneg = _symbol_position_stats(sym, dates, eod, min_d, max_d)
        ws.write(i + 1, base_stat + 0, float(mn))
        ws.write(i + 1, base_stat + 1, float(mx))
        ws.write(i + 1, base_stat + 2, _xls_stat_cell(mpos))
        ws.write(i + 1, base_stat + 3, _xls_stat_cell(mdpos))
        ws.write(i + 1, base_stat + 4, _xls_stat_cell(mneg))
        ws.write(i + 1, base_stat + 5, _xls_stat_cell(mdneg))

    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb.save(path)


def write_matrix_xlsx(
    path: str,
    dates: List[date],
    symbols: List[str],
    eod: Dict[Tuple[str, date], float],
    min_d: date,
    max_d: date,
) -> None:
    rows: List[Dict[str, object]] = []
    extreme_append: List[Tuple[str, float]] = []
    for sym in symbols:
        row: Dict[str, object] = {"symbol": sym}
        for d in dates:
            row[d.isoformat()] = _cell_net_usdt(sym, d, eod, min_d, max_d)
        mn, mx, mpos, mdpos, mneg, mdneg = _symbol_position_stats(sym, dates, eod, min_d, max_d)
        row[POSITION_STAT_COLS[0]] = mn
        row[POSITION_STAT_COLS[1]] = mx
        row[POSITION_STAT_COLS[2]] = mpos
        row[POSITION_STAT_COLS[3]] = mdpos
        row[POSITION_STAT_COLS[4]] = mneg
        row[POSITION_STAT_COLS[5]] = mdneg
        rows.append(row)
        mx_mean = _max_pos_mean_or_abs_neg_mean(mpos, mneg)
        if mx_mean > EXTREME_MEAN_THRESHOLD_U:
            extreme_append.append((sym, mx_mean))
    extreme_append.sort(key=lambda t: (-t[1], t[0]))
    out_df = pd.DataFrame(rows)
    cols = ["symbol"] + [d.isoformat() for d in dates] + list(POSITION_STAT_COLS)
    out_df = out_df[[c for c in cols if c in out_df.columns]]
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    out_df.to_excel(path, index=False, engine="openpyxl")
    _append_xlsx_extreme_mean_block(path, extreme_append)


def plot_symbol_series(
    charts_dir: str,
    symbol: str,
    dates: List[date],
    values: List[float],
) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt

    os.makedirs(charts_dir, exist_ok=True)
    xs = [datetime(d.year, d.month, d.day) for d in dates]
    arr = np.array(values, dtype=float)
    mask = np.isfinite(arr)
    vmean_pos: Optional[float]
    vmed_pos: Optional[float]
    vmean_neg: Optional[float]
    vmed_neg: Optional[float]
    if not mask.any():
        vmax = vmin = 0.0
        vmean_pos = vmed_pos = vmean_neg = vmed_neg = None
    else:
        sub = arr[mask]
        vmax = float(np.max(sub))
        vmin = float(np.min(sub))
        pos = sub[sub > 0.0]
        neg = sub[sub < 0.0]
        vmean_pos = float(np.mean(pos)) if pos.size else None
        vmed_pos = float(np.median(pos)) if pos.size else None
        vmean_neg = float(np.mean(neg)) if neg.size else None
        vmed_neg = float(np.median(neg)) if neg.size else None

    fig, ax = plt.subplots(figsize=(10, 4))
    ax.plot(xs, arr, label="net_usdt", color="C0", linewidth=1.0)
    ax.axhline(vmax, color="C1", linestyle="--", linewidth=0.9, label="max={:.6g}".format(vmax))
    ax.axhline(vmin, color="C2", linestyle="--", linewidth=0.9, label="min={:.6g}".format(vmin))
    if vmean_pos is not None:
        ax.axhline(
            vmean_pos,
            color="C3",
            linestyle=":",
            linewidth=0.9,
            label="mean(>0)={:.6g}".format(vmean_pos),
        )
    if vmed_pos is not None:
        ax.axhline(
            vmed_pos,
            color="C4",
            linestyle=":",
            linewidth=0.9,
            label="median(>0)={:.6g}".format(vmed_pos),
        )
    if vmean_neg is not None:
        ax.axhline(
            vmean_neg,
            color="C5",
            linestyle=":",
            linewidth=0.9,
            label="mean(<0)={:.6g}".format(vmean_neg),
        )
    if vmed_neg is not None:
        ax.axhline(
            vmed_neg,
            color="C6",
            linestyle=":",
            linewidth=0.9,
            label="median(<0)={:.6g}".format(vmed_neg),
        )
    ax.set_title(symbol)
    ax.legend(loc="best", fontsize=8)
    ax.grid(True, alpha=0.3)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
    fig.autofmt_xdate()
    fig.tight_layout()

    fp = os.path.join(charts_dir, safe_chart_basename(symbol))
    fig.savefig(fp, dpi=120)
    plt.close(fig)


def main() -> int:
    args = parse_args()
    start_col_d = datetime.strptime(args.start_col_date.strip(), "%Y-%m-%d").date()

    xls_p = (args.output_xls or "").strip()
    xlsx_p = (args.output_xlsx or "").strip()
    charts_dir = (args.charts_dir or "").strip()

    if not xls_p and not xlsx_p:
        sys.stderr.write("请指定 --output-xls 或 --output-xlsx（若二者都传空则无法导出）\n")
        return 2

    paths = list_pkl_paths(args.data_dir)
    raw = load_concat_pkls(paths)
    if raw.shape[0] == 0:
        sys.stderr.write("无数据：目录内无有效 PKL 或为空: {}\n".format(args.data_dir))
        return 1

    legs = ensure_user_id(raw)
    if not args.no_exclude:
        ex = get_excluded_user_ids_sync()
        uid = legs["user_id"].astype(str)
        legs = legs[~uid.isin(ex)].copy()
    legs = legs.sort_values("ts_text", kind="mergesort").reset_index(drop=True)

    sys.stderr.write(
        "载入行数={} 过滤后行数={} PKL文件数={}\n".format(raw.shape[0], legs.shape[0], len(paths))
    )

    if legs.shape[0] == 0:
        sys.stderr.write("过滤后无成交腿（可能全部被做市/模拟金排除），退出\n")
        return 1

    dates, symbols, eod, skipped, min_d, max_d = simulate_eod_matrix(legs, start_col_d)
    if skipped:
        sys.stderr.write("跳过未知 buy_sell 行数: {}\n".format(skipped))

    if not dates or not symbols:
        sys.stderr.write("无输出：过滤后无有效成交或日期列为空\n")
        return 1

    ncols = 1 + len(dates) + len(POSITION_STAT_COLS)
    if ncols > XLS_MAX_COLS and xls_p:
        sys.stderr.write(
            "列数 {} 超过 .xls 上限 {}，请改用 --output-xlsx 或缩短日期范围\n".format(
                ncols, XLS_MAX_COLS
            )
        )
        return 1

    if xls_p:
        write_matrix_xls(os.path.abspath(xls_p), dates, symbols, eod, min_d, max_d)
        sys.stderr.write("已写 xls: {}\n".format(os.path.abspath(xls_p)))

    if xlsx_p:
        write_matrix_xlsx(os.path.abspath(xlsx_p), dates, symbols, eod, min_d, max_d)
        sys.stderr.write("已写 xlsx: {}\n".format(os.path.abspath(xlsx_p)))

    if charts_dir:
        abs_charts = os.path.abspath(charts_dir)
        for sym in symbols:
            vals = [_cell_net_usdt(sym, d, eod, min_d, max_d) for d in dates]
            plot_symbol_series(abs_charts, sym, dates, vals)
        sys.stderr.write("已写图表目录: {} ({} 个 symbol)\n".format(abs_charts, len(symbols)))

    return 0


if __name__ == "__main__":
    sys.exit(main())
